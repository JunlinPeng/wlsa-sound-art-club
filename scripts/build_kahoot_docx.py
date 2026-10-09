#!/usr/bin/env python3
"""Build a Word (.docx) version of the Kahoot questions from the markdown source.

Reads src/meeting/kahoot-questions.md and writes a formatted .docx to the
project root. The correct answer for each question is bold + orange with a check.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src" / "meeting" / "kahoot-questions.md"
OUT = ROOT / "WLSA_Sound_Art_Club_Kahoot_Questions.docx"

INK = RGBColor(0x21, 0x1A, 0x12)
MUTED = RGBColor(0x8A, 0x7F, 0x70)
ORANGE = RGBColor(0xE8, 0x59, 0x0C)
BODY = RGBColor(0x4A, 0x42, 0x38)

text = SRC.read_text(encoding="utf-8")

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Segoe UI"
normal.font.size = Pt(11)

# Title
h = doc.add_paragraph()
hr = h.add_run("Kahoot — Sound Art Club Warm-up")
hr.bold = True
hr.font.size = Pt(20)
hr.font.color.rgb = INK

sub = doc.add_paragraph()
sr = sub.add_run("22 multiple-choice questions about the club and what we do.\n"
                 "Correct answer is in orange with a ✓.")
sr.italic = True
sr.font.size = Pt(11)
sr.font.color.rgb = MUTED

# Split into questions on the markdown "## " headings
parts = re.split(r"\n## ", text)
for part in parts[1:]:
    lines = part.split("\n")
    qtitle = lines[0].strip()

    qp = doc.add_paragraph()
    qp.paragraph_format.space_before = Pt(12)
    qp.paragraph_format.space_after = Pt(2)
    qr = qp.add_run(qtitle)
    qr.bold = True
    qr.font.size = Pt(12)
    qr.font.color.rgb = INK

    for ln in lines[1:]:
        m = re.match(r"^- ([A-D]\) .*)", ln)
        if not m:
            continue
        opt = m.group(1)
        is_correct = "**✓**" in ln or "**✓**" in ln
        opt = opt.replace("**✓**", "").replace("**✓**", "").strip()

        op = doc.add_paragraph()
        op.paragraph_format.left_indent = Inches(0.35)
        op.paragraph_format.space_after = Pt(1)
        orun = op.add_run(opt + ("  ✓" if is_correct else ""))
        orun.font.size = Pt(11)
        if is_correct:
            orun.bold = True
            orun.font.color.rgb = ORANGE
        else:
            orun.font.color.rgb = INK

doc.save(OUT)
print("Saved:", OUT)
