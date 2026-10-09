#!/usr/bin/env python3
"""Build the WLSA Sound Art Club first-meeting deck (matches the recruitment deck).

Opens the existing recruitment deck as a template (inherits its theme, fonts and
16:9 size), drops its slides, and builds four new slides in the same editorial
style: uppercase kicker -> big Georgia title -> Segoe UI body, with the club's
brand palette (orange #E8590C / blue #0E86C8 / coral #E04028, cream bg #FBF6EE).
The guess-the-sound slide links to src/game/index.html via a "Launch game" button.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = r'C:\Users\Junlin Peng\Downloads\WLSA_Sound_Art_Club.pptx'
OUT = r'D:\WLSA\CCA Setup\WLSA_Sound_Art_Club_First_Meeting.pptx'
PREVIEW = r'D:\WLSA\CCA Setup\.tmp\game-home.png'
GAME_URL = 'file:///D:/WLSA/CCA%20Setup/src/game/index.html'

# Cartoon characters from the previous deck (cream background, blends with slides)
MEDIA = r'D:\WLSA\CCA Setup\src\meeting\media'
SKYLINE = MEDIA + r'\skyline.png'
CHILD = MEDIA + r'\child.png'
NOTE_ORANGE = MEDIA + r'\note-orange.png'
NOTE_BLUE = MEDIA + r'\note-blue.png'
NOTE_CORAL = MEDIA + r'\note-coral.png'

# Brand palette (taken from the existing deck's direct formatting)
INK = RGBColor(0x21, 0x1A, 0x12)
MUTED = RGBColor(0x8A, 0x7F, 0x70)
BODY = RGBColor(0x4A, 0x42, 0x38)
ORANGE = RGBColor(0xE8, 0x59, 0x0C)
BLUE = RGBColor(0x0E, 0x86, 0xC8)
CORAL = RGBColor(0xE0, 0x40, 0x28)
CREAM = RGBColor(0xFB, 0xF6, 0xEE)
CREAM_TEXT = RGBColor(0xFF, 0xF9, 0xF2)

DISPLAY = 'Georgia'   # titles (matches existing deck)
SANS = 'Segoe UI'     # labels / body

prs = Presentation(SRC)
prs.slide_width = 12192000    # 13.333 in
prs.slide_height = 6858000    # 7.5 in

# --- Drop the existing slides (keep the theme/master) ---------------------
R_ID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
lst = prs.slides._sldIdLst
for sld in list(lst):
    prs.part.drop_rel(sld.get(R_ID))
    lst.remove(sld)

blank = next((l for l in prs.slide_layouts if l.name == 'Blank'), prs.slide_layouts[0])


def new_slide():
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = CREAM
    return s


def box(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def run(p, text, font, size, color, bold=False, italic=False):
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    return r


def kicker_left(slide, text):
    _, tf = box(slide, 1.25, 0.57, 9.0, 0.4)
    run(tf.paragraphs[0], text, SANS, 14, MUTED, bold=True)


def kicker_center(slide, text):
    _, tf = box(slide, 0, 0.62, 13.33, 0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, text, SANS, 14, MUTED, bold=True)


def title_center(slide, runs, size=66, top=1.15):
    _, tf = box(slide, 0, top, 13.33, 1.6)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for text, color in runs:
        run(p, text, DISPLAY, size, color, bold=True)


def title_left(slide, runs, size=40, top=1.15):
    _, tf = box(slide, 0.75, top, 11.8, 1.2)
    for text, color in runs:
        run(tf.paragraphs[0], text, DISPLAY, size, color, bold=True)


def body(slide, text, top, size=18, color=BODY, left=0.75, width=11.8):
    _, tf = box(slide, left, top, width, 1.0)
    run(tf.paragraphs[0], text, SANS, size, color)


def bullets(slide, items, top=2.35, size=19, left=0.75, width=8.4):
    _, tf = box(slide, left, top, width, 4.4)
    for i, (color, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        run(p, '\u25CF  ', SANS, size - 3, color, bold=True)
        run(p, text, SANS, size, INK)


def pic(slide, path, l, t, w):
    slide.shapes.add_picture(path, Inches(l), Inches(t), width=Inches(w))


def button(slide, text, l, t, w, h, fill=ORANGE, size=20, url=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, text, SANS, size, CREAM_TEXT, bold=True)
    if url:
        try:
            shp.click_action.hyperlink.address = url
        except Exception:
            p.runs[0].hyperlink.address = url
    return shp


# ===================== Slide 1 — Welcome ==================================
s = new_slide()
kicker_center(s, 'WELCOME  \u00B7  FIRST MEETING')
title_center(s, [('WLSA ', INK), ('Sound Art', ORANGE), (' Club', INK)], 66, 1.15)
_, tf = box(s, 0, 2.75, 13.33, 0.6)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, 'Hear the city.  ', DISPLAY, 26, BODY)
run(p, 'Save its memory.', DISPLAY, 26, ORANGE)
_, tf = box(s, 0, 3.5, 13.33, 0.4)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, 'A community-service club recording the disappearing sounds of our hometown.', SANS, 16, MUTED)

# cartoon decorations (skyline + music notes, matching the original title slide)
pic(s, SKYLINE, 0, 5.0, 13.33)
pic(s, NOTE_ORANGE, 1.7, 3.15, 0.7)
pic(s, NOTE_BLUE, 10.9, 3.05, 0.7)
pic(s, NOTE_CORAL, 11.7, 2.15, 0.6)

# ===================== Slide 2 — Introduce yourself =======================
s = new_slide()
kicker_left(s, "LET'S GET TO KNOW EACH OTHER")
title_left(s, [('Introduce ', INK), ('yourself.', BLUE)], 40)
bullets(s, [
    (ORANGE, 'Your name & grade'),
    (BLUE, 'One city sound you\u2019d love to save \u2014 and why'),
    (CORAL, 'Why you\u2019re curious about Sound Art'),
], top=2.4, size=19)
pic(s, CHILD, 9.2, 2.5, 2.9)

# ===================== Slide 3 — Kahoot ===================================
s = new_slide()
kicker_left(s, 'ACTIVITY 01  \u00B7  WARM-UP')
title_left(s, [("Let's play ", INK), ('Kahoot!', ORANGE)], 40)
body(s, 'Take out your device and go to kahoot.it', top=2.2, size=20)
_, tf = box(s, 0.75, 3.1, 8.0, 1.1)
run(tf.paragraphs[0], 'kahoot.it', DISPLAY, 48, BLUE, bold=True)
_, tf = box(s, 0.75, 4.6, 8.0, 0.6)
run(tf.paragraphs[0], 'Game PIN:  ', SANS, 22, MUTED, bold=True)
run(tf.paragraphs[0], '____________', SANS, 22, INK, bold=True)
pic(s, NOTE_BLUE, 11.4, 1.0, 0.7)

# ===================== Slide 4 — Sound Blind Box game =====================
s = new_slide()
kicker_left(s, 'ACTIVITY 02  \u00B7  GUESS THE SOUND')
title_left(s, [('Sound ', INK), ('Blind Box', ORANGE), ('.', INK)], 40)
body(s, 'Hear a mystery city sound \u2014 guess what it is.', top=1.9, size=18)
s.shapes.add_picture(PREVIEW, Inches(3.42), Inches(2.55), width=Inches(6.5))
button(s, 'Launch game  \u2192', 5.17, 6.45, 3.0, 0.8, ORANGE, 20, GAME_URL)
pic(s, NOTE_ORANGE, 11.5, 0.8, 0.7)
pic(s, NOTE_CORAL, 0.6, 5.6, 0.6)

prs.save(OUT)
print('Saved:', OUT)
print('slide count:', len(prs.slides._sldIdLst))
